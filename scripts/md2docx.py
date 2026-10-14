#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
md2docx.py —— 把论文 Markdown 转成结构化的 .docx（依赖 python-docx）。

支持：
  # / ## / ###  -> Heading 1/2/3
  - item        -> List Bullet
  1. item       -> List Number
  > quote       -> 缩进段落
  | a | b |     -> 表格（第二行 --- 为分隔）
  ```mermaid    -> 线性流程框（带边框的段落 + ↓ 箭头）
  ![alt](path)  -> 嵌入图片（下方附题注）
  **bold** `code` *italic* -> 行内格式

用法：
  PYTHONPATH=/home/yzm/.cache/pydocx python md2docx.py 论文.md 论文.docx
"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def set_east_asian_font(doc, latin="Times New Roman", east="宋体"):
    st = doc.styles["Normal"]
    st.font.name = latin
    st.font.size = Pt(10.5)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    rf.set(qn("w:eastAsia"), east)
    rf.set(qn("w:ascii"), latin)
    rf.set(qn("w:hAnsi"), latin)


def add_runs(par, text):
    """解析行内 **bold** / `code` / *italic* 并写入 runs。"""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
        elif tok.startswith("*"):
            par.add_run(tok[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def add_shaded_paragraph(doc, text, fill="EDF3FB", italic=False, indent=None):
    """带底纹/边框的段落（用于 mermaid 流程框）。"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if indent:
        p.paragraph_format.left_indent = Inches(indent)
    # 底纹
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)
    # 边框
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "4")
        e.set(qn("w:color"), "B9C4D2")
        pbdr.append(e)
    pPr.append(pbdr)
    parts = text.split("<br/>")
    for i, seg in enumerate(parts):
        if i:
            p.add_run().add_break()
        r = p.add_run(seg)
        r.italic = italic
    return p


def render_mermaid(doc, block_lines):
    """把 mermaid graph 线性化为 节点框 + ↓ 序列。"""
    nodes = []
    for ln in block_lines:
        for m in re.finditer(r"([A-Za-z0-9_]+)\[(.*?)\]", ln):
            label = m.group(2).replace("\\", "")
            if (m.group(1), label) not in nodes:
                nodes.append((m.group(1), label))
    for i, (_, label) in enumerate(nodes):
        if i:
            ap = doc.add_paragraph()
            ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ap.add_run("↓")
        add_shaded_paragraph(doc, label)


def parse_table(lines, i):
    """从 lines[i] 开始解析一张 markdown 表格，返回 (rows, next_i)。"""
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            rows.append(cells)
        i += 1
    return rows, i


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncol)
    try:
        t.style = "Light Grid Accent 1"
    except Exception:
        t.style = "Table Grid"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci in range(ncol):
            txt = row[ci] if ci < len(row) else ""
            par = cells[ci].paragraphs[0]
            add_runs(par, txt)
            if ri == 0:
                for r in par.runs:
                    r.bold = True
    doc.add_paragraph()


def main():
    src, dst = sys.argv[1], sys.argv[2]
    root = os.path.dirname(os.path.abspath(src))
    with open(src, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()
    set_east_asian_font(doc)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()

        # 空行
        if not s:
            i += 1
            continue

        # 代码块（仅处理 mermaid；其它原样作为等宽段落）
        if s.startswith("```"):
            lang = s[3:].strip().lower()
            i += 1
            block = []
            while i < n and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # 跳过结尾 ```
            if lang == "mermaid":
                render_mermaid(doc, block)
            else:
                for b in block:
                    p = doc.add_paragraph()
                    r = p.add_run(b)
                    r.font.name = "Consolas"
                    r.font.size = Pt(9)
            continue

        # 分隔线
        if re.fullmatch(r"-{3,}", s):
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            level = len(m.group(1))
            p = doc.add_paragraph(style=f"Heading {min(level,4)}")
            add_runs(p, m.group(2))
            i += 1
            continue

        # 图片
        m = re.match(r"^!\[(.*?)\]\((.*?)\)$", s)
        if m:
            alt, path = m.group(1), m.group(2)
            ap = os.path.join(root, path)
            if os.path.exists(ap):
                try:
                    doc.add_picture(ap, width=Inches(6.0))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cr = cap.add_run(alt)
                    cr.italic = True
                    cr.font.size = Pt(9)
                    cr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                except Exception as e:
                    doc.add_paragraph(f"[图片缺失: {path} ({e})]")
            else:
                doc.add_paragraph(f"[图片缺失: {path}]")
            i += 1
            continue

        # 表格
        if s.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        # 引用
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, s.lstrip(">").strip())
            for r in p.runs:
                r.italic = True
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\d+\.\s+(.*)$", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(1))
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, s)
        i += 1

    doc.save(dst)
    print(f"已生成: {dst}")


if __name__ == "__main__":
    main()
